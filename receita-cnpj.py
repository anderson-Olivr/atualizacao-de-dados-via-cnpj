import re
import time
import shutil
import requests
from pathlib import Path
from datetime import datetime
from openpyxl import load_workbook

# =========================
# VARIÁVEIS IMPORTANTES
# =========================

NOME_PLANILHA = "plan-teste.xlsx"
CAMINHO_PLANILHA = Path("planilha-a-analisar") / NOME_PLANILHA

BACKUP_A_CADA = 20
PASTA_BACKUP = Path("bkp-planilha")

# =========================
# COLUNAS
# =========================

COLUNAS = [
    "CNPJ",
    "Razão Social",
    "Fantasia",
    "Endereço",
    "Número",
    "Complemento",
    "Bairro",
    "Cidade",
    "Estado",
    "CEP",
    "Telefone",
    "Email",
    "Contato",
    "IE",
    "Simples Nacional",
    "Situação Cadastral",
    "Atualizado"
]

API_URL = "https://www.receitaws.com.br/v1/cnpj/{}"


def limpar_cnpj(cnpj):
    return re.sub(r"\D", "", str(cnpj or ""))


def cnpj_valido(cnpj):
    cnpj = limpar_cnpj(cnpj)

    if len(cnpj) != 14:
        return False

    if cnpj == cnpj[0] * 14:
        return False

    def calcular_digito(cnpj_parcial, pesos):
        soma = sum(int(d) * p for d, p in zip(cnpj_parcial, pesos))
        resto = soma % 11
        return "0" if resto < 2 else str(11 - resto)

    digito1 = calcular_digito(cnpj[:12], [5, 4, 3, 2, 9, 8, 7, 6, 5, 4, 3, 2])
    digito2 = calcular_digito(cnpj[:13], [6, 5, 4, 3, 2, 9, 8, 7, 6, 5, 4, 3, 2])

    return cnpj[-2:] == digito1 + digito2


def proximo_nome_backup():
    PASTA_BACKUP.mkdir(exist_ok=True)

    data = datetime.now().strftime("%Y-%m-%d")
    nome_base = Path(NOME_PLANILHA).stem

    existentes = list(PASTA_BACKUP.glob(f"{nome_base}-{data}-*.xlsx"))

    numeros = []
    for arquivo in existentes:
        try:
            numero = int(arquivo.stem.split("-")[-1])
            numeros.append(numero)
        except ValueError:
            pass

    proximo = max(numeros, default=0) + 1
    return PASTA_BACKUP / f"{nome_base}-{data}-{proximo}.xlsx"


def salvar_backup():
    destino = proximo_nome_backup()
    shutil.copy2(CAMINHO_PLANILHA, destino)
    print(f"Backup salvo: {destino}")


def consultar_receitaws(cnpj):
    while True:
        try:
            resposta = requests.get(API_URL.format(cnpj), timeout=20)

            if resposta.status_code == 429:
                print("Limite de consultas atingido. Aguardando 1 minuto...")
                time.sleep(60)
                continue

            dados = resposta.json()

            if dados.get("status") == "ERROR":
                mensagem = dados.get("message", "").lower()

                if "limite" in mensagem or "too many" in mensagem:
                    print("Limite da ReceitaWS atingido. Aguardando 1 minuto...")
                    time.sleep(60)
                    continue

                return dados

            return dados

        except requests.exceptions.RequestException:
            print("rede instavel ou desconectado")
            time.sleep(60)


def traduzir_situacao(situacao):
    situacao = str(situacao or "").strip().upper()

    if situacao == "ATIVA":
        return "Ativa"

    return "Desativada"


def traduzir_simples(dados):
    simples = dados.get("simples")

    if isinstance(simples, dict):
        optante = simples.get("optante")

        if optante is True:
            return "Optante"

        if optante is False:
            return "Não Optante"

    if str(simples).lower() in ["true", "sim", "optante"]:
        return "Optante"

    if str(simples).lower() in ["false", "nao", "não", "não optante"]:
        return "Não Optante"

    return ""


def atualizar_planilha():
    wb = load_workbook(CAMINHO_PLANILHA)
    ws = wb.worksheets[0]

    for col, titulo in enumerate(COLUNAS, start=1):
        ws.cell(row=1, column=col).value = titulo

    atualizados = 0

    for linha in range(2, ws.max_row + 1):
        cnpj_original = ws.cell(row=linha, column=1).value

        atualizado = str(
            ws.cell(row=linha, column=17).value or ""
        ).strip().lower()

        if atualizado == "ok":
            continue

        cnpj = limpar_cnpj(cnpj_original)

        if not cnpj:
            continue

        if not cnpj_valido(cnpj):
            print(f"Linha {linha}: CNPJ inválido ou incompleto: {cnpj_original}")
            continue

        print(f"Consultando linha {linha}: {cnpj}")

        dados = consultar_receitaws(cnpj)

        if dados.get("status") == "ERROR":
            print(f"Linha {linha}: erro na consulta - {dados.get('message')}")
            continue

        telefone = str(dados.get("telefone", "") or "").replace(",", " / ")

        ws.cell(row=linha, column=1).value = cnpj
        ws.cell(row=linha, column=2).value = dados.get("nome", "")
        ws.cell(row=linha, column=3).value = dados.get("fantasia", "")
        ws.cell(row=linha, column=4).value = dados.get("logradouro", "")
        ws.cell(row=linha, column=5).value = dados.get("numero", "")
        ws.cell(row=linha, column=6).value = dados.get("complemento", "")
        ws.cell(row=linha, column=7).value = dados.get("bairro", "")
        ws.cell(row=linha, column=8).value = dados.get("municipio", "")
        ws.cell(row=linha, column=9).value = dados.get("uf", "")
        ws.cell(row=linha, column=10).value = dados.get("cep", "")

        # Telefone
        ws.cell(row=linha, column=11).value = telefone

        # Email
        ws.cell(row=linha, column=12).value = dados.get("email", "")

        # Contato
        qsa = dados.get("qsa", [])

        if qsa and len(qsa) > 0:
            contato = qsa[0].get("nome", "")
        else:
            contato = ""

        ws.cell(row=linha, column=13).value = contato

        # IE
        ws.cell(row=linha, column=14).value = dados.get("ie", "")

        # Simples Nacional
        ws.cell(row=linha, column=15).value = traduzir_simples(dados)

        # Situação Cadastral
        ws.cell(row=linha, column=16).value = traduzir_situacao(
            dados.get("situacao", "")
        )

        # Atualizado
        ws.cell(row=linha, column=17).value = "ok"

        atualizados += 1

        if atualizados % BACKUP_A_CADA == 0:
            wb.save(CAMINHO_PLANILHA)
            salvar_backup()

        time.sleep(21)

    nome_final = Path(NOME_PLANILHA).stem + "-ATUALIZADA.xlsx"
    caminho_final = Path("planilha-a-analisar") / nome_final

    wb.save(caminho_final)
    print(f"Planilha atualizada salva em: {caminho_final}")


if __name__ == "__main__":
    atualizar_planilha()